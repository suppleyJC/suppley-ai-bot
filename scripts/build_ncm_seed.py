#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gerador do seed de alíquotas de Imposto de Importação (II) por NCM, a partir do
arquivo OFICIAL e CONSOLIDADO do MDIC:

    "Anexos I a X da Resolução GECEX 272/21" (Tarifas Vigentes)
    https://www.gov.br/mdic/.../arquivos-listas/<data>-anexos-i-a-x-resolucao-gecex-272-21.xlsx

Por que este arquivo: ele consolida a TEC base (Anexo I), a alíquota efetivamente
aplicada pelo Brasil (Anexo II) e as ELEVAÇÕES TEMPORÁRIAS (Anexo IX - DCC, ex.: aço
a 25% via Res. GECEX 740). É a mesma base que o Simulador da Receita usa.

II EFETIVA (prioridade):
  1. Anexo IX (DCC) vigente HOJE  -> elevação temporária (ex.: 25%)
  2. Anexo II "Alíquota aplicada" -> alíquota estrutural aplicada pelo Brasil
  3. Anexo I  "TEC (%)"           -> tarifa externa comum base
  4. default por capítulo         -> último recurso (raro)

Saída: data/ncm_import.sql.gz com INSERT ... ON DUPLICATE KEY UPDATE que ATUALIZA
APENAS iiRate + notes (preserva descrição, IPI, PIS, COFINS já cadastrados).

USO (no servidor, onde está o xlsx):
    python3 scripts/build_ncm_seed.py /tmp/tec_vigente.xlsx
    # opcional: caminho de saída
    python3 scripts/build_ncm_seed.py /tmp/tec_vigente.xlsx data/ncm_import.sql.gz
"""
import sys
import os
import re
import gzip
import datetime
import unicodedata

try:
    import openpyxl
except ImportError:
    sys.exit("ERRO: openpyxl não instalado. Rode: apt-get install -y python3-openpyxl")

# Defaults de II por capítulo (fração em basis points: 14% = 1400) — só p/ fallback.
DEFAULT_II_BY_CHAPTER = {
    "01": 400, "02": 1000, "03": 1000, "04": 1600, "05": 600, "06": 600, "07": 1000,
    "08": 1000, "09": 1000, "10": 800, "11": 1200, "12": 800, "13": 1400, "14": 800,
    "15": 1000, "16": 1600, "17": 1600, "18": 1400, "19": 1600, "20": 1600, "21": 1600,
    "22": 2000, "23": 800, "24": 2000, "25": 400, "26": 400, "27": 0, "28": 1200,
    "29": 1200, "30": 800, "31": 600, "32": 1400, "33": 1800, "34": 1400, "35": 1400,
    "36": 1800, "37": 1400, "38": 1400, "39": 1400, "40": 1400, "41": 1000, "42": 2000,
    "43": 2000, "44": 1000, "45": 1200, "46": 1800, "47": 600, "48": 1400, "49": 0,
    "50": 1400, "51": 1200, "52": 1800, "53": 1200, "54": 1800, "55": 1800, "56": 1800,
    "57": 3500, "58": 2600, "59": 1800, "60": 1800, "61": 3500, "62": 3500, "63": 3500,
    "64": 3500, "65": 2000, "66": 2000, "67": 2000, "68": 1000, "69": 1200, "70": 1200,
    "71": 1800, "72": 1200, "73": 1400, "74": 1000, "75": 800, "76": 1200, "78": 1000,
    "79": 1000, "80": 1000, "81": 800, "82": 1800, "83": 1800, "84": 1400, "85": 1600,
    "86": 1400, "87": 3500, "88": 0, "89": 1400, "90": 1400, "91": 2000, "92": 2000,
    "93": 2000, "94": 1800, "95": 2000, "96": 1800, "97": 400, "98": 0, "99": 0,
}

NCM_RE = re.compile(r"^\d{4}\.\d{2}\.\d{2}$")


def norm(s):
    """minúsculas, sem acento, sem ordinais (º/ª), espaços colapsados — p/ casar cabeçalhos."""
    if s is None:
        return ""
    s = str(s)
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("º", "").replace("ª", "").replace("°", "")
    s = re.sub(r"\s+", " ", s)
    return s.lower().strip()


def clean_ncm(raw):
    """'7308.40.00' -> '73084000' (8 dígitos) ou None."""
    if raw is None:
        return None
    digits = re.sub(r"\D", "", str(raw))
    return digits[:8] if len(digits) >= 8 else None


def parse_rate_to_bp(raw):
    """'14' | '12.6' | '12,6' | '25' -> basis points (12,6% = 1260). None se inválido."""
    if raw is None:
        return None
    s = str(raw).replace("%", "").replace(",", ".").strip()
    if s == "" or s == "-" or s.upper() == "NT":
        return None
    try:
        pct = float(s)
    except ValueError:
        return None
    if pct < 0:
        return None
    return round(pct * 100)


def find_header_row(ws, aliases_by_key, required, max_scan=12):
    """
    Acha a linha de cabeçalho e devolve (idx_linha, {chave: idx_coluna}).
    aliases_by_key: dict {chave: [aliases normalizados]}.
    required: lista de chaves obrigatórias (as demais são best-effort).
    """
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= max_scan:
            break
        normd = [norm(c) for c in row]
        colmap = {}
        for key, aliases in aliases_by_key.items():
            for ci, val in enumerate(normd):
                if any(val == a or val.startswith(a) for a in aliases):
                    colmap.setdefault(key, ci)
                    break
        if all(k in colmap for k in required):
            return i, colmap
    return None, {}


def today():
    return datetime.date.today()


def to_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if v is None:
        return None
    s = str(v).strip()
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r"(\d{2})/(\d{2})/(\d{4})", s)
    if m:
        return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None


def main():
    if len(sys.argv) < 2:
        sys.exit("USO: python3 scripts/build_ncm_seed.py <arquivo.xlsx> [saida.sql.gz]")
    xlsx = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "..", "data", "ncm_import.sql.gz")
    out = os.path.abspath(out)
    if not os.path.exists(xlsx):
        sys.exit(f"ERRO: não encontrei {xlsx}")

    file_tag = os.path.basename(xlsx)
    gen_date = today()
    print(f"[seed] Lendo {xlsx} ...")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    # ---- Anexo I - TEC (base): NCM -> bp ----
    tec = {}
    ws1 = next((w for w in wb.worksheets if "anexo i" in norm(w.title) and "tec" in norm(w.title)), None)
    if ws1 is not None:
        for row in ws1.iter_rows(values_only=True):
            if not row:
                continue
            code = row[0]
            if code is None or not NCM_RE.match(str(code).strip()):
                continue
            ncm = clean_ncm(code)
            if not ncm:
                continue
            # taxa = primeiro número plausível a partir da col 2
            rate = None
            for c in row[2:]:
                rate = parse_rate_to_bp(c)
                if rate is not None:
                    break
            if rate is not None:
                tec[ncm] = rate
    print(f"[seed] Anexo I (TEC): {len(tec)} NCMs")

    # ---- Anexo II - Alíquota aplicada: NCM -> (bp, desc) ----
    aplicada = {}
    desc_map = {}
    ws2 = next((w for w in wb.worksheets if "anexo ii" in norm(w.title)), None)
    if ws2 is not None:
        hdr_i, cols = find_header_row(ws2, {
            "ncm": ["ncm"],
            "aplicada": ["aliquota aplicada", "aliquota aplicada (%)"],
            "desc": ["descricao"],
        }, required=["ncm", "aplicada"])
        if hdr_i is not None and "aplicada" in cols:
            for i, row in enumerate(ws2.iter_rows(values_only=True)):
                if i <= hdr_i:
                    continue
                ncm = clean_ncm(row[cols["ncm"]]) if cols["ncm"] < len(row) else None
                if not ncm:
                    continue
                rate = parse_rate_to_bp(row[cols["aplicada"]]) if cols["aplicada"] < len(row) else None
                if rate is not None:
                    aplicada[ncm] = rate
                if "desc" in cols and cols["desc"] < len(row) and row[cols["desc"]]:
                    desc_map[ncm] = str(row[cols["desc"]]).strip()
        else:
            print("[seed] AVISO: não achei coluna 'Alíquota aplicada' no Anexo II")
    print(f"[seed] Anexo II (aplicada): {len(aplicada)} NCMs")

    # ---- Anexo IX - DCC (elevações temporárias vigentes): NCM -> (bp, termino) ----
    dcc = {}
    ws9 = next((w for w in wb.worksheets if "anexo ix" in norm(w.title) or "dcc" in norm(w.title)), None)
    if ws9 is not None:
        hdr_i, cols = find_header_row(ws9, {
            "ncm": ["ncm"],
            "aliq": ["aliquota", "aliquota (%)"],
            "ini": ["inicio de vigencia", "inicio"],
            "fim": ["termino de vigencia", "termino"],
            "ex": ["n ex", "no ex", "numero ex", "ex"],
        }, required=["ncm", "aliq"])
    if ws9 is not None and hdr_i is not None:
        for i, row in enumerate(ws9.iter_rows(values_only=True)):
            if i <= hdr_i:
                continue
            ncm = clean_ncm(row[cols["ncm"]]) if cols.get("ncm", 99) < len(row) else None
            if not ncm:
                continue
            # só aplica a NCM inteira (sem EX específico)
            if "ex" in cols and cols["ex"] < len(row):
                exv = row[cols["ex"]]
                if exv not in (None, "", "-"):
                    continue
            rate = parse_rate_to_bp(row[cols["aliq"]]) if cols.get("aliq", 99) < len(row) else None
            if rate is None:
                continue
            ini = to_date(row[cols["ini"]]) if cols.get("ini", 99) < len(row) else None
            fim = to_date(row[cols["fim"]]) if cols.get("fim", 99) < len(row) else None
            # vigente hoje?
            if ini and gen_date < ini:
                continue
            if fim and gen_date > fim:
                continue
            # mantém a maior elevação se houver duplicidade
            prev = dcc.get(ncm)
            if prev is None or rate > prev[0]:
                dcc[ncm] = (rate, fim)
    print(f"[seed] Anexo IX (DCC vigente em {gen_date}): {len(dcc)} NCMs")

    # ---- Consolida II efetiva por NCM ----
    all_ncms = set(tec) | set(aplicada) | set(dcc)
    rows = []
    n_dcc = n_apl = n_tec = n_def = 0
    for ncm in sorted(all_ncms):
        if ncm in dcc:
            bp, fim = dcc[ncm]
            origem = f"II {bp/100:g}% DCC vigente ate {fim} (Res.GECEX 740/elevacao) | MDIC {gen_date}"
            n_dcc += 1
        elif ncm in aplicada:
            bp = aplicada[ncm]
            origem = f"II {bp/100:g}% aplicada (Anexo II) | MDIC {gen_date}"
            n_apl += 1
        elif ncm in tec:
            bp = tec[ncm]
            origem = f"II {bp/100:g}% TEC (Anexo I) | MDIC {gen_date}"
            n_tec += 1
        else:
            bp = DEFAULT_II_BY_CHAPTER.get(ncm[:2], 1400)
            origem = f"II {bp/100:g}% estimada por capitulo | MDIC {gen_date}"
            n_def += 1
        desc = desc_map.get(ncm, "")
        rows.append((ncm, desc, bp, origem))

    print(f"[seed] Consolidado: {len(rows)} NCMs "
          f"(DCC={n_dcc}, aplicada={n_apl}, TEC={n_tec}, default={n_def})")

    # ---- Verificação rápida do caso 7308.40.00 ----
    for ncm, desc, bp, origem in rows:
        if ncm == "73084000":
            print(f"[seed] >>> 73084000: II={bp/100:g}%  ({origem})")
            break

    # ---- Gera SQL (atualiza só iiRate + notes; preserva IPI/PIS/COFINS/descrição) ----
    def esc(s):
        return str(s).replace("\\", "\\\\").replace("'", "''")

    parts = ["SET NAMES utf8mb4;", "SET autocommit=0;"]
    CHUNK = 1000
    for start in range(0, len(rows), CHUNK):
        chunk = rows[start:start + CHUNK]
        parts.append(
            "INSERT INTO ncm_tax_rates "
            "(ncmCode,description,iiRate,ipiRate,pisRate,cofinsRate,mercosulIiRate,notes) VALUES")
        vals = []
        for ncm, desc, bp, origem in chunk:
            vals.append(
                f"('{ncm}','{esc(desc)}',{bp},0,210,1025,0,'{esc(origem)}')")
        parts.append(",\n".join(vals) +
                     "\nON DUPLICATE KEY UPDATE iiRate=VALUES(iiRate),notes=VALUES(notes);")
    parts.append("COMMIT;")
    sql = "\n".join(parts) + "\n"

    os.makedirs(os.path.dirname(out), exist_ok=True)
    with gzip.open(out, "wt", encoding="utf-8") as f:
        f.write(sql)
    size_kb = os.path.getsize(out) / 1024
    print(f"[seed] Gravado {out} ({size_kb:.0f} KB)")
    print(f"[seed] Fonte: {file_tag}")
    print("[seed] OK. Agora rode:  bash scripts/load-ncm.sh")


if __name__ == "__main__":
    main()
